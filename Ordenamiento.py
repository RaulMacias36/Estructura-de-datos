import subprocess
import sys
import os

# --- BLOQUE DE AUTO-INSTALACIÓN ---
def install_dependencies():
    try:
        import openpyxl
    except ImportError:
        print("Instalando librería para Excel (openpyxl)...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])

install_dependencies()

# --- IMPORTACIONES ---
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import time
import json
import random
import csv
from datetime import datetime
from openpyxl import Workbook, load_workbook

class ExternalSortingApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Procesador Universal de Ordenamiento Externo")
        self.root.geometry("1000x900")
        self.root.configure(bg="#0f172a")

        self.all_data = []
        self.canvas_width = 920
        self.canvas_height = 250
        self.output_dir = "archivos_generados"
        if not os.path.exists(self.output_dir): os.makedirs(self.output_dir)
            
        self.setup_ui()

    def setup_ui(self):
        # Título
        header = tk.Label(self.root, text="LABORATORIO DE FUSIÓN MULTIFORMATO", 
                         bg="#0f172a", fg="#10b981", font=("Segoe UI", 22, "bold"))
        header.pack(pady=10)

        # --- PANEL: GENERADOR COMPLETO ---
        gen_frame = tk.LabelFrame(self.root, text=" ⚡ Generador Automático (TXT, JSON, CSV, DAT, LOG, EXCEL) ", 
                                 bg="#1e293b", fg="#f8fafc", font=("Segoe UI", 10, "bold"), padx=15, pady=10)
        gen_frame.pack(fill="x", padx=25, pady=5)

        tk.Button(gen_frame, text="📊 Crear Excel (H/V/R)", command=self.generate_excel_auto,
                  bg="#10b981", fg="white", relief="flat", font=("Segoe UI", 8, "bold")).pack(side="left", padx=5)
        
        for fmt, col in [("TXT", "#ef4444"), ("JSON", "#f59e0b"), ("CSV", "#3b82f6"), ("DAT", "#6366f1"), ("LOG", "#8b5cf6")]:
            tk.Button(gen_frame, text=fmt, command=lambda f=fmt.lower(): self.generate_simple_auto(f),
                      bg=col, fg="white", relief="flat", width=6, font=("Arial", 8, "bold")).pack(side="left", padx=2)

        # --- PANEL: CONFIGURACIÓN DE EXCEL ---
        opt_frame = tk.LabelFrame(self.root, text=" ⚙️ Configuración de Fusión de Hojas Excel ", 
                                 bg="#334155", fg="white", font=("Segoe UI", 10, "bold"), padx=15, pady=12)
        opt_frame.pack(fill="x", padx=25, pady=10)

        self.excel_mode = tk.StringVar(value="todo")
        tk.Radiobutton(opt_frame, text="Todo el Libro", variable=self.excel_mode, value="todo",
                       bg="#334155", fg="white", font=("Arial", 9), selectcolor="#0f172a").pack(side="left")
        tk.Radiobutton(opt_frame, text="Índices de hojas:", variable=self.excel_mode, value="custom",
                       bg="#334155", fg="white", font=("Arial", 9), selectcolor="#0f172a").pack(side="left", padx=10)
        
        self.custom_sheets = ttk.Entry(opt_frame, width=30)
        self.custom_sheets.insert(0, "1, 2, 3") 
        self.custom_sheets.pack(side="left", padx=5)
        tk.Label(opt_frame, text="(Ej: 1, 3 para Hoja 1 y 3)", bg="#334155", fg="#94a3b8", font=("Arial", 8)).pack(side="left")

        tk.Button(opt_frame, text="📁 Cargar y Combinar", command=self.load_multiple_files,
                  bg="#38bdf8", fg="#0f172a", font=("Segoe UI", 10, "bold"), relief="flat").pack(side="right")

        # --- ÁREA GRÁFICA ---
        self.canvas = tk.Canvas(self.root, width=self.canvas_width, height=self.canvas_height, 
                                bg="#020617", highlightthickness=1, highlightbackground="#1e293b")
        self.canvas.pack(pady=10)

        # --- PANEL: ALGORITMOS ---
        btn_frame = tk.Frame(self.root, bg="#0f172a")
        btn_frame.pack(pady=10)

        for text, cmd, col in [("Intercalación", self.run_intercalacion, "#22c55e"),
                               ("Mezcla Directa", self.run_directa, "#3b82f6"),
                               ("Mezcla Equilibrada", self.run_equilibrada, "#a855f7")]:
            tk.Button(btn_frame, text=text, command=cmd, width=18, bg=col, fg="white", 
                      font=("Segoe UI", 10, "bold"), relief="flat").pack(side="left", padx=10)

        self.status_var = tk.StringVar(value="Listo.")
        tk.Label(self.root, textvariable=self.status_var, bg="#1e293b", fg="#38bdf8", anchor="w", padx=10).pack(side="bottom", fill="x")

    # --- GENERADORES ---
    def generate_excel_auto(self):
        try:
            wb = Workbook()
            ws1 = wb.active; ws1.title = "Horizontal"
            for c in range(1, 21): ws1.cell(1, c, random.randint(1, 100))
            ws2 = wb.create_sheet("Vertical")
            for r in range(1, 21): ws2.cell(r, 1, random.randint(1, 100))
            ws3 = wb.create_sheet("Aleatoria")
            for _ in range(20): ws3.cell(random.randint(1, 10), random.randint(1, 10), random.randint(1, 100))
            path = self._get_auto_path("xlsx")
            wb.save(path)
            messagebox.showinfo("Generador", f"Excel creado: {path}")
        except Exception as e: messagebox.showerror("Error", str(e))

    def generate_simple_auto(self, fmt):
        nums = [random.randint(1, 1000) for _ in range(50)]
        path = self._get_auto_path(fmt)
        if fmt == "json":
            with open(path, 'w') as f: json.dump({"data": nums}, f)
        elif fmt == "csv":
            with open(path, 'w', newline='') as f: csv.writer(f).writerow(nums)
        else:
            with open(path, 'w') as f: f.write(" ".join(map(str, nums)))
        messagebox.showinfo("Generador", f"Archivo {fmt.upper()} creado.")

    def _get_auto_path(self, ext):
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        return os.path.join(self.output_dir, f"gen_{ts}.{ext}")

    # --- CARGA Y FUSIÓN ---
    def load_multiple_files(self):
        file_paths = filedialog.askopenfilenames()
        if not file_paths: return
        self.all_data = []
        
        # Obtener lista de índices numéricos (1-based para el usuario)
        indices_usuario = []
        if self.excel_mode.get() == "custom":
            try:
                indices_usuario = [int(s.strip()) for s in self.custom_sheets.get().split(",") if s.strip().isdigit()]
            except:
                messagebox.showwarning("Aviso", "Formato de índices no válido. Se usará Hoja 1.")
                indices_usuario = [1]

        try:
            for path in file_paths:
                ext = os.path.splitext(path)[1].lower()
                if ext == ".xlsx":
                    wb = load_workbook(path, data_only=True)
                    nombres_hojas = wb.sheetnames
                    
                    if self.excel_mode.get() == "todo":
                        target_sheets = nombres_hojas
                    else:
                        # Convertir a 0-based y validar existencia
                        target_sheets = []
                        for idx in indices_usuario:
                            zero_idx = idx - 1
                            if 0 <= zero_idx < len(nombres_hojas):
                                target_sheets.append(nombres_hojas[zero_idx])
                    
                    for sn in target_sheets:
                        for row in wb[sn].iter_rows(values_only=True):
                            for cell in row:
                                if isinstance(cell, (int, float)):
                                    self.all_data.append(int(cell))

                elif ext == ".json":
                    with open(path, 'r') as f: self.all_data.extend(self.extract_nums(json.load(f)))
                elif ext == ".csv":
                    with open(path, 'r') as f:
                        for r in csv.reader(f): self.all_data.extend([int(x) for x in r if x.strip().isdigit()])
                else: 
                    with open(path, 'r') as f:
                        text = f.read().replace(',', ' ').split()
                        self.all_data.extend([int(x) for x in text if x.lstrip('-').isdigit()])
            
            self.draw_bars(self.all_data)
            self.status_var.set(f"Cargados {len(self.all_data)} elementos de {len(file_paths)} archivo(s).")
        except Exception as e: 
            messagebox.showerror("Error", f"Error al cargar: {e}")

    def extract_nums(self, obj):
        nums = []
        if isinstance(obj, list):
            for i in obj: nums.extend(self.extract_nums(i))
        elif isinstance(obj, dict):
            for v in obj.values(): nums.extend(self.extract_nums(v))
        elif isinstance(obj, (int, float)): nums.append(int(obj))
        return nums

    # --- LÓGICA DE ORDENAMIENTO ---
    def draw_bars(self, arr, highlights={}):
        self.canvas.delete("all")
        if not arr: return
        n = len(arr); w = self.canvas_width / n; mx = max(arr) if arr else 1
        for i, v in enumerate(arr):
            color = highlights.get(i, "#38bdf8")
            h = (v / mx) * (self.canvas_height - 10)
            self.canvas.create_rectangle(i*w, self.canvas_height-h, (i+1)*w, self.canvas_height, fill=color, outline="#020617")
        self.root.update_idletasks()

    def _animate_merge(self, arr, start, mid, end):
        left = arr[start:mid]; right = arr[mid:end]; i = j = 0
        for k in range(start, end):
            if i < len(left) and (j >= len(right) or left[i] <= right[j]):
                arr[k] = left[i]; i += 1
            else: arr[k] = right[j]; j += 1
            if len(arr) <= 150: 
                self.draw_bars(arr, {k: "#f43f5e"})
                time.sleep(0.01) # Velocidad original recuperada

    def run_directa(self):
        if not self.all_data: return
        data = list(self.all_data); n = len(data); width = 1
        while width < n:
            for i in range(0, n, width * 2): self._animate_merge(data, i, min(i+width, n), min(i+2*width, n))
            width *= 2
        self.final_step(data)

    def run_intercalacion(self):
        if not self.all_data: return
        data = list(self.all_data); mid = len(data) // 2
        data[:mid], data[mid:] = sorted(data[:mid]), sorted(data[mid:])
        self._animate_merge(data, 0, mid, len(data))
        self.final_step(data)

    def run_equilibrada(self):
        if not self.all_data: return
        data = list(self.all_data)
        def get_runs(d):
            runs = []; curr = [0]
            for i in range(1, len(d)):
                if d[i] < d[i-1]: runs.append(curr); curr = [i]
                else: curr.append(i)
            runs.append(curr); return runs
        while True:
            runs = get_runs(data)
            if len(runs) <= 1: break
            for i in range(0, len(runs)-1, 2): self._animate_merge(data, runs[i][0], runs[i+1][0], runs[i+1][-1] + 1)
        self.final_step(data)

    def final_step(self, data):
        self.draw_bars(data)
        if messagebox.askyesno("Guardar", "¿Exportar resultado a Excel?"):
            path = filedialog.asksaveasfilename(defaultextension=".xlsx")
            if path:
                wb = Workbook(); ws = wb.active; ws.title = "Ordenados"
                for i, v in enumerate(data): ws.cell(i+1, 1, v)
                wb.save(path); messagebox.showinfo("Éxito", "Resultado guardado correctamente.")

if __name__ == "__main__":
    root = tk.Tk(); app = ExternalSortingApp(root); root.mainloop()